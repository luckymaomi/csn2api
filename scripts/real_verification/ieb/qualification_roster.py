from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import requests
from openpyxl import load_workbook

from ..app_http import post_json, require_completed, serve_app, verify_frontend_route, wait_for_final
from ..common import ROOT, result_summary


APP_ROOT = ROOT / "public/tool/app/http-qualification-roster-helper"
EXPECTED_PERSONNEL_COLUMNS = [
    "员工号", "姓名", "技术基地", "注册基地", "运行基地", "技术信息", "实际机型", "保护等级",
    "性别", "在职状态", "出生日期", "国籍", "政治面貌", "所属部门", "行政职务", "行政级别",
]


def verify_api(verifier: Any, executor: Any) -> None:
    filters = {
        "fleetCdToArray": ["777"],
        "activeStatusArray": ["ZAIZHI", "WAIBU", "TUIXIU", "LIZHI"],
        "operQualArray": verifier.input("iebRosterQualification"),
    }
    for endpoint, params in (
        ("qualification-roster.filter-page", {}),
        ("qualification-roster.operating-qualifications", {}),
        ("qualification-roster.query", filters),
        ("qualification-roster.native-export", filters),
    ):
        verifier.call(
            "ieb",
            f"api.{endpoint}",
            lambda endpoint=endpoint, params=params: executor.execute(endpoint, params),
            detail=lambda value: f"HTTP {value.get('status')}，{result_summary(value)}" if isinstance(value, dict) else str(value),
        )


def verify_app(verifier: Any) -> None:
    from http_qualification_roster.manager import RunManager
    from http_qualification_roster.server import create_server

    try:
        qualification = verifier.inputs.get("iebRosterQualification", "").strip()
        if not qualification:
            raise ValueError("本地验收参数缺少 iebRosterQualification")
        with TemporaryDirectory(prefix="csn2api-real-roster-") as directory:
            manager = RunManager(Path(directory))
            server = create_server(APP_ROOT, 0, manager)
            with serve_app(server, manager) as base_url:
                verify_frontend_route(base_url, "/api/run")
                post_json(base_url, "/api/session/verify", {"credentials": verifier.credentials["ieb"]})
                metadata = post_json(base_url, "/api/metadata/load", {})
                post_json(
                    base_url,
                    "/api/run",
                    {
                        "filters": {
                            "fleetCdToArray": ["777"],
                            "activeStatusArray": ["ZAIZHI", "WAIBU", "TUIXIU", "LIZHI"],
                        },
                        "qualificationCodes": [qualification],
                    },
                )
                state = require_completed(wait_for_final(base_url))
                columns = state.get("resultColumns")
                if not isinstance(columns, list) or any(not isinstance(column, str) for column in columns):
                    raise RuntimeError("APP 未返回动态人员字段")
                missing = [column for column in EXPECTED_PERSONNEL_COLUMNS if column not in columns]
                if missing or "#" in columns or "" in columns:
                    raise RuntimeError(f"APP 人员字段与门户表格不一致，缺少：{', '.join(missing) or '无'}")
                results = state.get("results")
                if not isinstance(results, list) or not results:
                    raise RuntimeError("APP 未返回人员资质明细")
                for index, row in enumerate(results, start=1):
                    if not isinstance(row, dict) or row.get("sequence") != index:
                        raise RuntimeError("APP 人员资质明细序号不连续")
                    if list((row.get("fields") or {}).keys()) != columns or row.get("qualificationCode") != qualification:
                        raise RuntimeError("APP 人员字段或资质类别与本轮查询不一致")
                response = requests.get(f"{base_url}/api/download/excel", timeout=30)
                if response.status_code != 200:
                    raise RuntimeError(f"APP Excel 下载返回 HTTP {response.status_code}")
                workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
                try:
                    rows = list(workbook.active.iter_rows(values_only=True))
                finally:
                    workbook.close()
                expected_headers = ["序号", *columns, "资质类别"]
                actual_headers = [str(value or "") for value in rows[0]] if rows else []
                if actual_headers != expected_headers:
                    raise RuntimeError("APP 页面与 Excel 列结构不一致")
                if len(rows) - 1 != len(results):
                    raise RuntimeError("APP 页面与 Excel 明细行数不一致")
                if any(row[0] != index or row[-1] != qualification for index, row in enumerate(rows[1:], start=1)):
                    raise RuntimeError("APP Excel 序号或资质类别不正确")
        progress = state.get("progress", {})
        verifier.record("ieb", "app.qualification-roster.metadata", "PASS", f"动态筛选字段 {len(metadata.get('filterFields', []))}，动态资格 {len(metadata.get('qualifications', []))}")
        verifier.record("ieb", "app.qualification-roster.query", "PASS", f"前端 /api/run 已接通，人员字段 {len(columns)} 项，关系 {progress.get('relationships', 0)} 条")
        verifier.record("ieb", "app.qualification-roster.native-export", "PASS", "页面与临时 Excel 的完整人员字段、连续序号和资质类别一致，验收文件已清理")
    except Exception as error:  # noqa: BLE001
        verifier.fail("ieb", "app.qualification-roster", error)
