from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import requests
from openpyxl import Workbook, load_workbook
from ieb_protocol.flight_history_contract import is_summary_label

from ..app_http import post_json, require_completed, serve_app, verify_frontend_route, wait_for_final
from ..common import ROOT


APP_ROOT = ROOT / 'public/tool/app/http-flight-history-helper'


def _check(verifier: Any, headers: list[str], rows: list[dict[str, str]]) -> None:
    summaries = [row for row in rows if is_summary_label(row.get('日期', ''))]
    if '飞行时间' not in headers or len(summaries) != 1:
        raise RuntimeError('未返回唯一汇总行')
    expected = verifier.inputs.get('flightHistoryExpectedTime', '')
    if expected and summaries[0].get('飞行时间') != expected:
        raise RuntimeError('汇总飞行时间与本地验收参数不一致')
    verifier.flight_history_reference = (headers, rows)


def verify_api(verifier: Any, executor: Any) -> None:
    def query():
        result = executor.execute('flight-detail.query-history', {
            'staffNum': verifier.input('iebStaffNumber'), 'staffName': verifier.input('iebStaffName'),
        })
        table = result['data']['tables'][0]
        _check(verifier, table['columns'], table['rows'])
        return result
    verifier.call('ieb', 'api.flight-history.query', query, detail=lambda value: f"人员归属、全页 {len(value['data']['tables'][0]['rows'])} 行、原表头与预期时间核对通过")


def verify_app(verifier: Any) -> None:
    from http_flight_history.manager import RunManager
    from http_flight_history.server import create_server
    import base64

    try:
        staff, name = verifier.input('iebStaffNumber'), verifier.input('iebStaffName')
        with TemporaryDirectory(prefix='csn2api-real-flight-history-') as directory:
            manager = RunManager(Path(directory))
            server = create_server(APP_ROOT, 0, manager)
            address = server.server_address
            with serve_app(server, manager) as base_url:
                verify_frontend_route(base_url, '/api/run')
                post_json(base_url, '/api/session/verify', {'credentials': verifier.credentials['ieb']})
                workbook = Workbook()
                workbook.active.append(['员工号', '姓名'])
                workbook.active.append([staff, name])
                buffer = BytesIO()
                workbook.save(buffer)
                workbook.close()
                inputs = [
                    {'inputMode': 'paste', 'pastedText': f'员工号\t姓名\n{staff}\t{name}'},
                    {'inputMode': 'excel', 'excelBase64': base64.b64encode(buffer.getvalue()).decode('ascii')},
                ]
                for payload in inputs:
                    post_json(base_url, '/api/check-data', payload)
                    post_json(base_url, '/api/run', payload)
                    state = require_completed(wait_for_final(base_url))
                    result = state['results'][0]
                    reference = getattr(verifier, 'flight_history_reference', None)
                    _check(verifier, result['headers'], result['rows'])
                    if reference and (result['headers'], result['rows']) != reference:
                        raise RuntimeError('API 与 APP 全页结果不一致')
                    response = requests.get(f'{base_url}/api/download/original', timeout=30)
                    response.raise_for_status()
                    exported = load_workbook(BytesIO(response.content), data_only=True)
                    try:
                        rows = list(exported.worksheets[0].values)
                        if len(rows) != len(result['rows']) + 1 or list(rows[0]) != ['员工号', '姓名', *result['headers']]:
                            raise RuntimeError('Excel 全页表头或行数异常')
                        expected_rows = [[staff, name, *[row[h] or None for h in result['headers']]] for row in result['rows']]
                        if [list(row) for row in rows[1:]] != expected_rows:
                            raise RuntimeError('Excel 与页面全页数据不一致')
                    finally:
                        exported.close()
            import socket
            with socket.socket() as sock:
                if sock.connect_ex(address) == 0:
                    raise RuntimeError('临时 APP 端口未释放')
        verifier.record('ieb', 'app.flight-history.query', 'PASS', 'Excel/粘贴入口、完整全页数据、导出与 API 一致；临时服务端口已释放')
    except Exception as error:
        verifier.fail('ieb', 'app.flight-history.query', error)
