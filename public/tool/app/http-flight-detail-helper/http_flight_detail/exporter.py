from __future__ import annotations

import json
import os
import uuid
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook

from ieb_protocol.flight_detail_contract import DETAIL_HEADERS

from .models import MonthOutcome, QueryRecord


DETAIL_OUTPUT_HEADERS = ["序号", "姓名", "员工号", *DETAIL_HEADERS[1:]]
STATUS_HEADERS = ["序号", "姓名", "员工号", "查询月份", "状态", "错误说明", "航班数量"]


def _save_atomic(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temp)
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def _append_jsonl(path: Path, rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")))
            stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())


def _read_jsonl(path: Path) -> Iterable[dict[str, object]]:
    if not path.is_file():
        return []

    def rows() -> Iterable[dict[str, object]]:
        with path.open("r", encoding="utf-8") as stream:
            for line in stream:
                try:
                    value = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if isinstance(value, dict):
                    yield value

    return rows()


class ResultExporter:
    """追加式检查点写入器，最终按需封装为 Excel。"""

    def __init__(self, output_directory: Path, run_id: str):
        self.output_directory = output_directory
        self.path = output_directory / f"HTTP航班明细_{run_id}.xlsx"
        self._details_log = output_directory / "detail-rows.jsonl"
        self._statuses_log = output_directory / "status-rows.jsonl"
        self._errors_log = output_directory / "input-errors.jsonl"
        self._detail_index = 1
        self._written_outcomes = 0
        self._outcome_keys: set[tuple[int, str]] = set()
        self._errors_written = False
        self._finalized = False
        self._closed = False

    def add_input_errors(self, input_errors: list[str]) -> None:
        if self._errors_written or not input_errors:
            return
        _append_jsonl(self._errors_log, ({"index": index, "message": message} for index, message in enumerate(input_errors, 1)))
        self._errors_written = True

    def append_batch(self, outcomes: list[MonthOutcome]) -> None:
        pending_details: list[dict[str, object]] = []
        pending_statuses: list[dict[str, object]] = []
        for outcome in outcomes[self._written_outcomes:]:
            self._outcome_keys.add((outcome.record_index, outcome.month))
            for row in outcome.rows or []:
                pending_details.append({
                    "序号": self._detail_index,
                    "姓名": outcome.record.name,
                    "员工号": outcome.record.employee_id,
                    **{header: row.get(header, "") for header in DETAIL_HEADERS[1:]},
                })
                self._detail_index += 1
            pending_statuses.append({
                "序号": outcome.record_index + 1,
                "姓名": outcome.record.name,
                "员工号": outcome.record.employee_id,
                "查询月份": outcome.month,
                "状态": outcome.status,
                "错误说明": outcome.error,
                "航班数量": len(outcome.rows or []),
            })
        if pending_details:
            _append_jsonl(self._details_log, pending_details)
        if pending_statuses:
            _append_jsonl(self._statuses_log, pending_statuses)
        self._written_outcomes = len(outcomes)

    def flush(self) -> Path:
        """检查点数据已在追加写入时 fsync；此方法保留统一生命周期接口。"""
        self.output_directory.mkdir(parents=True, exist_ok=True)
        return self.path

    def finalize(self, records: list[QueryRecord] | None = None, input_errors: list[str] | None = None) -> Path:
        if self._finalized:
            return self.path
        self.add_input_errors(input_errors or [])
        if records:
            missing = []
            for record_index, record in enumerate(records):
                for month in record.months():
                    if (record_index, month) not in self._outcome_keys:
                        missing.append({
                            "序号": record_index + 1,
                            "姓名": record.name,
                            "员工号": record.employee_id,
                            "查询月份": month,
                            "状态": "未查询",
                            "错误说明": "尚未执行",
                            "航班数量": 0,
                        })
                        self._outcome_keys.add((record_index, month))
            if missing:
                _append_jsonl(self._statuses_log, missing)
        workbook = Workbook(write_only=True)
        details = workbook.create_sheet("航班明细")
        details.append(DETAIL_OUTPUT_HEADERS)
        for row in _read_jsonl(self._details_log):
            details.append([row.get(header, "") for header in DETAIL_OUTPUT_HEADERS])
        statuses = workbook.create_sheet("查询状态")
        statuses.append(STATUS_HEADERS)
        for row in _read_jsonl(self._statuses_log):
            statuses.append([row.get(header, "") for header in STATUS_HEADERS])
        if self._errors_log.is_file():
            errors = workbook.create_sheet("输入错误")
            errors.append(["序号", "错误说明"])
            for row in _read_jsonl(self._errors_log):
                errors.append([row.get("index", ""), row.get("message", "")])
        try:
            _save_atomic(workbook, self.path)
        finally:
            workbook.close()
        self._finalized = True
        return self.path

    def close(self) -> None:
        if self._finalized:
            for path in (self._details_log, self._statuses_log, self._errors_log):
                path.unlink(missing_ok=True)
        self._closed = True

    def write(self, records: list[QueryRecord], input_errors: list[str], outcomes: list[MonthOutcome]) -> Path:
        self.add_input_errors(input_errors)
        self.append_batch(outcomes)
        return self.finalize(records, input_errors)
