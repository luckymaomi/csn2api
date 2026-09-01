from __future__ import annotations

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ieb_protocol.flight_detail_contract import month_range

from .models import InputPayload, QueryRecord


EXPECTED_HEADERS = ["员工号", "姓名", "开始年月", "结束年月"]


def normalize_month(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    match = re.fullmatch(r"\s*(\d{4})/(\d{1,2})\s*", str(value or ""))
    if not match:
        raise ValueError("月份格式必须为 YYYY/M 或 YYYY/MM")
    year, month = int(match.group(1)), int(match.group(2))
    if not 1 <= month <= 12:
        raise ValueError("月份必须是有效的年月")
    return f"{year:04d}-{month:02d}"


def normalize_employee_id(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text.isdigit() or len(text) > 6:
        return None
    return text.zfill(6)


def _record(values: list[Any], source: str) -> QueryRecord:
    if len(values) != 4:
        raise ValueError("每行必须只有四列：员工号、姓名、开始年月、结束年月")
    employee_id = normalize_employee_id(values[0])
    if not employee_id:
        raise ValueError("员工号必须是 6 位数字")
    name = str(values[1] or "").strip()
    if not name:
        raise ValueError("姓名不能为空")
    start = normalize_month(values[2])
    end = normalize_month(values[3])
    month_range(start, end)
    return QueryRecord(employee_id, name, start, end, source)


def _validate_headers(headers: list[Any]) -> None:
    normalized = [str(value or "").strip() for value in headers]
    if normalized != EXPECTED_HEADERS:
        raise ValueError("表头必须严格为：员工号、姓名、开始年月、结束年月")


def _parse_rows(rows: list[list[Any]], source_prefix: str) -> tuple[list[QueryRecord], list[str]]:
    records: list[QueryRecord] = []
    errors: list[str] = []
    seen: set[str] = set()
    for line_number, values in enumerate(rows, start=2):
        if not any(str(value or "").strip() for value in values):
            continue
        try:
            item = _record(values, f"{source_prefix}{line_number}行")
        except ValueError as error:
            errors.append(f"第{line_number}行：{error}")
            continue
        if item.employee_id in seen:
            errors.append(f"第{line_number}行：重复员工号，已跳过")
            continue
        seen.add(item.employee_id)
        records.append(item)
    return records, errors


def read_excel_records(source: str | Path | bytes) -> tuple[list[QueryRecord], list[str]]:
    workbook_source = BytesIO(source) if isinstance(source, bytes) else source
    workbook = load_workbook(workbook_source, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Excel 没有数据")
    _validate_headers(list(rows[0]))
    return _parse_rows([list(row) for row in rows[1:]], "Excel 第")


def _delimiter(line: str) -> str:
    return "\t" if "\t" in line else ","


def parse_pasted_records(text: str) -> tuple[list[QueryRecord], list[str]]:
    lines = [line for line in str(text or "").splitlines() if line.strip()]
    if not lines:
        raise ValueError("请粘贴包含表头和数据的四列表格")
    delimiter = _delimiter(lines[0])
    rows = list(csv.reader(StringIO("\n".join(lines)), delimiter=delimiter))
    _validate_headers(rows[0])
    return _parse_rows(rows[1:], "粘贴第")


def read_input(payload: InputPayload) -> tuple[list[QueryRecord], list[str]]:
    if payload.input_mode == "excel":
        if payload.excel_bytes is None:
            raise ValueError("请选择 Excel 文件")
        return read_excel_records(payload.excel_bytes)
    if payload.input_mode == "paste":
        return parse_pasted_records(payload.pasted_text)
    raise ValueError("不支持的输入方式")
