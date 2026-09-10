from __future__ import annotations

import os
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ieb_protocol.training_arrangement_models import TrainingWindow
from .area_defaults import area_details
from .models import AreaWindow


@dataclass(frozen=True)
class ExportPaths:
    excel: Path


def _safe(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "")).rstrip(" .")
    return cleaned or "run"


def _atomic_save(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


class ResultExporter:
    def __init__(self, output_directory: Path, run_id: str) -> None:
        self.paths = ExportPaths(
            Path(output_directory) / f"训练安排查询_{_safe(run_id)}.xlsx"
        )

    def write(self, windows: list[TrainingWindow]) -> ExportPaths:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "汇总"
        summary.append(["时间范围", "开始日期", "结束日期", "页数", "记录数", "门户总数"])
        for cell in summary[1]:
            cell.font = Font(bold=True)
        for window in windows:
            summary.append([
                window.label,
                window.start.isoformat(),
                window.end.isoformat(),
                window.pages,
                len(window.rows),
                window.total if window.total is not None else len(window.rows),
            ])
            sheet = workbook.create_sheet(_safe(window.label)[:31])
            sheet.append(window.headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in window.rows:
                sheet.append([row.get(header, "") for header in window.headers])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(42, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = summary.dimensions
        _atomic_save(workbook, self.paths.excel)
        workbook.close()
        return self.paths

    def write_area(self, windows: list[AreaWindow], summary_rows: list[dict[str, object]]) -> ExportPaths:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "分类汇总"
        summary.append(["类型", "本周", "", "本月", "", "年度", ""])
        summary.append(["", "人次", "通过率", "人次", "通过率", "人次", "通过率"])
        summary.merge_cells("A1:A2")
        summary.merge_cells("B1:C1")
        summary.merge_cells("D1:E1")
        summary.merge_cells("F1:G1")
        for row in summary.iter_rows(min_row=1, max_row=2):
            for cell in row:
                cell.font = Font(bold=True)
        for cell in summary[1]:
            cell.font = Font(bold=True)
        for item in summary_rows:
            values = [item.get(key, {}) for key in ("week", "month", "year")]
            formatted = []
            for value in values:
                if isinstance(value, dict):
                    formatted.extend([
                        value.get("count", 0),
                        self._rate_text(value.get("rate")),
                    ])
                else:
                    formatted.extend([0, "/"])
            summary.append([
                item.get("type", ""),
                *formatted,
            ])
        summary.freeze_panes = "B3"
        summary.auto_filter.ref = f"A2:G{summary.max_row}"
        detail = workbook.create_sheet("明细")
        detail_headers, detail_rows = area_details(windows)
        detail.append(detail_headers)
        for cell in detail[1]:
            cell.font = Font(bold=True)
        for row in detail_rows:
            detail.append([row.get(header, "") for header in detail_headers])
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions
        for sheet in (summary, detail):
            for index, column in enumerate(sheet.columns, start=1):
                width = min(42, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[get_column_letter(index)].width = width
        area_path = self.paths.excel.with_name(f"区域训练默认导出_{_safe(self.paths.excel.stem)}.xlsx")
        _atomic_save(workbook, area_path)
        workbook.close()
        return ExportPaths(area_path)

    @staticmethod
    def _rate_text(value: object) -> str:
        if value is None:
            return "/"
        number = float(value)
        return f"{int(number) if number.is_integer() else number}%"
