from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .flatten import flatten_detail, flatten_listing
from .models import PilotOutcome


LEADING_HEADERS = ["序号"]
TRAILING_HEADERS = ["来源行", "输入姓名", "输入档案号", "抓取状态", "失败原因"]
BASE_HEADERS = [*LEADING_HEADERS, *TRAILING_HEADERS]


def merge_export_headers(existing: Iterable[str], incoming: Iterable[str]) -> list[str]:
    fixed_headers = {*LEADING_HEADERS, *TRAILING_HEADERS}
    dynamic_headers: list[str] = []
    for header in [*existing, *incoming]:
        if header not in fixed_headers and header not in dynamic_headers:
            dynamic_headers.append(header)
    return [*LEADING_HEADERS, *dynamic_headers, *TRAILING_HEADERS]


def build_export_row(outcome: PilotOutcome) -> dict[str, str]:
    business_fields = flatten_listing(outcome.record.listing)
    if outcome.detail:
        business_fields.update(flatten_detail(outcome.detail))
    for header in BASE_HEADERS:
        business_fields.pop(header, None)
    return {
        "序号": str(outcome.record.index + 1),
        **business_fields,
        "来源行": outcome.record.source,
        "输入姓名": outcome.record.input_name,
        "输入档案号": outcome.record.input_archive_no,
        "抓取状态": outcome.status,
        "失败原因": outcome.error,
    }


class ResultExporter:
    def __init__(self, output_dir: Path, run_id: str) -> None:
        self.output_dir = output_dir
        self.run_id = run_id
        self.path = self.output_dir / f"caac-pilot-export-{self.run_id}.xlsx"
        self.temp_path = self.output_dir / f".{self.path.name}.tmp.xlsx"
        self.workbook = Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "飞行员信息"
        self.sheet.freeze_panes = "A2"
        self.headers = list(BASE_HEADERS)
        self.header_positions = {header: index for index, header in enumerate(self.headers, 1)}
        self._write_headers(self.headers)

    def _write_headers(self, headers: Iterable[str]) -> None:
        header_fill = PatternFill("solid", fgColor="D9E2F3")
        for header in headers:
            column = self.header_positions[header]
            cell = self.sheet.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _save_checkpoint(self) -> None:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.headers))}{max(1, self.sheet.max_row)}"
        self.workbook.save(self.temp_path)
        os.replace(self.temp_path, self.path)

    def _add_dynamic_headers(self, row_headers: Iterable[str]) -> list[str]:
        next_headers = merge_export_headers(self.headers, row_headers)
        new_headers = [header for header in next_headers if header not in self.header_positions]
        if not new_headers:
            return []
        first_trailing_column = self.header_positions[TRAILING_HEADERS[0]]
        self.sheet.insert_cols(first_trailing_column, amount=len(new_headers))
        self.headers = next_headers
        self.header_positions = {header: index for index, header in enumerate(self.headers, 1)}
        self._write_headers(self.headers)
        return new_headers

    def append(self, outcome: PilotOutcome) -> Path:
        row = build_export_row(outcome)
        new_headers = self._add_dynamic_headers(row)
        row_index = self.sheet.max_row + 1
        for header, value in row.items():
            cell = self.sheet.cell(row_index, self.header_positions[header], value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header == "抓取状态":
                cell.font = Font(color="008000" if value == "成功" else "C00000", bold=True)
        for header in new_headers:
            column = self.header_positions[header]
            width = max(12, min(60, len(header) + 2))
            self.sheet.column_dimensions[get_column_letter(column)].width = width
        self._save_checkpoint()
        return self.path

    def finalize(self) -> Path:
        self._save_checkpoint()
        return self.path

    def write(self, outcomes: Iterable[PilotOutcome]) -> Path:
        for outcome in outcomes:
            self.append(outcome)
        return self.finalize()
