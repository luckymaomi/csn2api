from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

from .models import InputIssue, InputPayload, PilotRecord


NAME_HEADER = "姓名"
ARCHIVE_HEADER = "档案号"


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def archive_key(value: Any) -> str:
    text = normalize_text(value)
    return (text.lstrip("0") or "0") if text.isdigit() else text


def read_excel_records(source: str | Path | BinaryIO) -> tuple[list[PilotRecord], list[InputIssue]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    records: list[PilotRecord] = []
    issues: list[InputIssue] = []
    seen: set[tuple[str, str]] = set()
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            return [], [InputIssue(0, "", "", "Excel 为空", "Excel")]
        headers = [normalize_text(value) for value in header_row]
        if NAME_HEADER not in headers:
            return [], [InputIssue(1, "", "", "未找到姓名表头", "Excel 第1行")]
        name_index = headers.index(NAME_HEADER)
        archive_index = headers.index(ARCHIVE_HEADER) if ARCHIVE_HEADER in headers else None
        for row_number, row in enumerate(rows, start=2):
            if not row or not any(normalize_text(value) for value in row):
                continue
            name = normalize_text(row[name_index] if name_index < len(row) else "")
            archive_no = normalize_text(row[archive_index] if archive_index is not None and archive_index < len(row) else "")
            source_label = f"Excel 第{row_number}行"
            if not name:
                issues.append(InputIssue(row_number, name, archive_no, "姓名不能为空", source_label))
                continue
            key = (name, archive_key(archive_no))
            if key in seen:
                issues.append(InputIssue(row_number, name, archive_no, "重复人员，已跳过重复请求", source_label))
                continue
            seen.add(key)
            records.append(PilotRecord(len(records), {}, input_name=name, input_archive_no=archive_no, source=source_label, input_row=row_number))
    finally:
        workbook.close()
    return records, issues


def read_input(payload: InputPayload) -> tuple[list[PilotRecord], list[InputIssue]]:
    if not payload.excel_bytes:
        raise ValueError("请选择 Excel 名单")
    if Path(payload.excel_name).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("只支持 .xlsx 或 .xlsm 文件")
    return read_excel_records(BytesIO(payload.excel_bytes))
