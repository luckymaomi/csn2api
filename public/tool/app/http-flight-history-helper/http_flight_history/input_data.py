from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .models import InputPayload, QueryRecord

EMPLOYEE_HEADERS = {'员工号', '工号', '员工编号'}
NAME_HEADERS = {'姓名', '员工姓名'}


def normalize_employee_id(value: object) -> str:
    text = str(value if value is not None else '').strip()
    if re.fullmatch(r'\d+\.0+', text):
        text = text.split('.', 1)[0]
    if not re.fullmatch(r'\d{1,6}', text):
        raise ValueError('员工号必须为六位数字')
    return text.zfill(6)


def _collect(rows: list[tuple[object, object, str]]) -> tuple[list[QueryRecord], list[str]]:
    records, errors, seen = [], [], set()
    for staff, name_value, source in rows:
        try:
            staff_number = normalize_employee_id(staff)
            name = str(name_value or '').strip()
            if not name or not re.fullmatch(r'[\w\u00b7\u2022\u30fb -]+', name) or name.isdigit():
                raise ValueError('姓名为空或格式无效')
            if staff_number in seen:
                raise ValueError('重复员工号，已跳过')
            seen.add(staff_number)
            records.append(QueryRecord(staff_number, name, source))
        except ValueError as error:
            errors.append(f'{source}：{error}')
    return records, errors


def parse_pasted_records(text: str) -> tuple[list[QueryRecord], list[str]]:
    rows = []
    for index, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        values = re.split(r'[\t,，;； ]+', line.strip(), maxsplit=1)
        if values[0] in EMPLOYEE_HEADERS:
            continue
        rows.append((values[0], values[1] if len(values) > 1 else '', f'粘贴第{index}行'))
    return _collect(rows)


def read_excel_records(source: str | Path | bytes) -> tuple[list[QueryRecord], list[str]]:
    workbook = load_workbook(BytesIO(source) if isinstance(source, bytes) else source, data_only=True, read_only=True)
    try:
        rows = iter(workbook.active.iter_rows(values_only=True))
        headers = [str(value or '').strip() for value in next(rows, ())]
        def locate(aliases: set[str]) -> int:
            matches = [i for i, value in enumerate(headers) if value in aliases]
            if len(matches) != 1:
                raise ValueError('Excel 必须包含唯一的员工号和姓名表头')
            return matches[0]
        staff_index, name_index = locate(EMPLOYEE_HEADERS), locate(NAME_HEADERS)
        values = []
        for index, row in enumerate(rows, 2):
            if any(value not in (None, '') for value in row):
                values.append((row[staff_index] if staff_index < len(row) else None,
                               row[name_index] if name_index < len(row) else None,
                               f'Excel 第{index}行'))
        return _collect(values)
    finally:
        workbook.close()


def read_input(payload: InputPayload) -> tuple[list[QueryRecord], list[str]]:
    if payload.input_mode == 'excel':
        if not payload.excel_bytes:
            raise ValueError('请选择 Excel 文件')
        return read_excel_records(payload.excel_bytes)
    if payload.input_mode == 'paste':
        if not payload.pasted_text.strip():
            raise ValueError('请粘贴员工号和姓名')
        return parse_pasted_records(payload.pasted_text)
    raise ValueError('不支持的输入方式')