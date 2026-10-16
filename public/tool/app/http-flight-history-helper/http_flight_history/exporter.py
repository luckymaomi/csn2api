from __future__ import annotations

import os
import uuid
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from .models import QueryOutcome, QueryRecord


@dataclass(frozen=True)
class OutputPaths:
    original: Path


class ResultExporter:
    """每名员工完成后原子保存；主表保留身份与时间轴全部数据行。"""

    def __init__(self, output_directory: Path, run_id: str):
        self.paths = OutputPaths(output_directory / f'飞行经历时间_{run_id}.xlsx')

    def write(self, records: list[QueryRecord], input_errors: list[str], outcomes: list[QueryOutcome]) -> OutputPaths:
        headers = list(dict.fromkeys(header for item in outcomes if item.result for header in item.result.headers))
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '飞行经历时间'
        sheet.append(['员工号', '姓名', *headers])
        for item in outcomes:
            if item.status != '成功' or item.result is None:
                continue
            for values in item.result.rows:
                sheet.append([item.record.employee_id, item.record.name, *[values.get(header, '') for header in headers]])
                for cell in sheet[sheet.max_row]:
                    cell.data_type = 's'
                    cell.number_format = '@'
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 16
        errors = workbook.create_sheet('处理记录')
        errors.append(['来源', '员工号', '姓名', '状态', '说明'])
        by_index = {item.index: item for item in outcomes}
        for index, record in enumerate(records):
            item = by_index.get(index)
            errors.append([record.source, record.employee_id, record.name, item.status if item else '未查询', item.error if item else ''])
        for error in input_errors:
            errors.append(['输入检查', '', '', '已跳过', error])
        for row in errors.iter_rows():
            for cell in row:
                cell.data_type = 's'
        path = self.paths.original
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(f'.{uuid.uuid4().hex}.tmp.xlsx')
        try:
            workbook.save(temporary)
            os.replace(temporary, path)
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
        return self.paths