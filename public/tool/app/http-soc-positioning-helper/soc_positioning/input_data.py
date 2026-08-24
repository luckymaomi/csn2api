"""Excel and pasted-text normalization for SOC positioning."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .constants import APPLY_REASON_CODES, APPLY_REASON_LABELS
from .models import InputPayload


HEADER_ALIASES = {
    "员工号": ("员工号", "工号", "员工编号"),
    "姓名": ("姓名", "中文名"),
    "航班日期": ("航班日期", "日期", "飞行日期"),
    "出发地": ("出发地", "出发站", "出发站1"),
    "目的地": ("目的地", "到达地", "到达站", "到达站1"),
    "航班号": ("航班号", "航班", "航班号1"),
    "申请原因": ("申请原因", "原因"),
    "备注": ("备注", "申请备注"),
    "是否需要录入": ("是否需要录入", "是否录入", "录入"),
}


def normalize_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalize_employee_id(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def normalize_date(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        compact = str(int(value)) if float(value).is_integer() else ""
        if re.fullmatch(r"(?:19|20)\d{6}", compact):
            value = compact
        else:
            try:
                return from_excel(value).strftime("%Y-%m-%d")
            except (TypeError, ValueError, OverflowError):
                return ""
    text = normalize_text(value)
    match = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", text)
    if match:
        text = "-".join(match.groups())
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if not match:
        return ""
    try:
        return date(*(int(part) for part in match.groups())).isoformat()
    except ValueError:
        return ""


def _header_indexes(header_row: tuple[object, ...]) -> tuple[dict[str, int], list[str]]:
    headers = [normalize_text(value) for value in header_row]
    indexes: dict[str, int] = {}
    missing: list[str] = []
    for field, aliases in HEADER_ALIASES.items():
        index = next((headers.index(alias) for alias in aliases if alias in headers), None)
        if index is None and field not in {"备注", "是否需要录入"}:
            missing.append(field)
        elif index is not None:
            indexes[field] = index
    return indexes, missing


def _value(row: tuple[object, ...], indexes: dict[str, int], field: str) -> object:
    index = indexes.get(field)
    return row[index] if index is not None and index < len(row) else None


def _yes(value: object) -> bool:
    return normalize_text(value).lower() in {"", "是", "是的", "yes", "y", "1", "true", "需要"}


def _finalize(raw: dict[str, object]) -> tuple[dict[str, object] | None, str]:
    employee_id = normalize_employee_id(raw.get("员工号"))
    if not re.fullmatch(r"\d{6}", employee_id):
        return None, f"员工号格式错误 [{employee_id}]"
    flight_date = normalize_date(raw.get("航班日期"))
    if not flight_date:
        return None, "航班日期格式错误"
    dep, arv = normalize_text(raw.get("出发地")), normalize_text(raw.get("目的地"))
    if not dep or not arv:
        return None, "出发地和目的地不能为空"
    reason = normalize_text(raw.get("申请原因"))
    code = reason if reason in APPLY_REASON_LABELS else APPLY_REASON_CODES.get(reason, "")
    if not code:
        return None, f"申请原因不在静态字典中 [{reason}]"
    remark = normalize_text(raw.get("备注"))
    if len(remark) > 24:
        return None, "备注不能超过24个字符"
    return {
        "员工号": employee_id, "姓名": normalize_text(raw.get("姓名")),
        "航班日期": flight_date, "出发地": dep, "目的地": arv,
        "航班号": normalize_text(raw.get("航班号")), "申请原因代码": code,
        "申请原因": APPLY_REASON_LABELS[code], "备注": remark,
        "是否需要录入": normalize_text(raw.get("是否需要录入")),
        "来源行": raw.get("来源行", ""),
    }, ""


def parse_excel(content: bytes) -> tuple[list[dict[str, object]], list[str]]:
    records: list[dict[str, object]] = []
    errors: list[str] = []
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        indexes, missing = _header_indexes(header_row)
        if missing:
            return [], [f"Excel缺少表头: {', '.join(missing)}"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and value != "" for value in row):
                continue
            raw = {field: _value(row, indexes, field) for field in HEADER_ALIASES}
            raw["来源行"] = row_number
            if not _yes(raw.get("是否需要录入")):
                continue
            record, problem = _finalize(raw)
            if problem:
                errors.append(f"第{row_number}行: {problem}")
            elif record:
                records.append(record)
    finally:
        workbook.close()
    return records, errors


def parse_paste(text: str) -> tuple[list[dict[str, object]], list[str]]:
    records: list[dict[str, object]] = []
    errors: list[str] = []
    for line_number, line in enumerate(str(text or "").splitlines(), start=1):
        if not line.strip():
            continue
        parts = [part.strip() for part in line.split("\t" if "\t" in line else "|")]
        if len(parts) < 7:
            errors.append(f"第{line_number}条: 至少需要7列（员工号、姓名、日期、出发地、目的地、航班号、申请原因）")
            continue
        raw = dict(zip(("员工号", "姓名", "航班日期", "出发地", "目的地", "航班号", "申请原因", "备注"), parts + [""]))
        raw["来源行"] = line_number
        record, problem = _finalize(raw)
        if problem:
            errors.append(f"第{line_number}条: {problem}")
        elif record:
            records.append(record)
    return records, errors


def read_input(payload: InputPayload) -> tuple[list[dict[str, object]], list[str]]:
    if payload.input_mode == "excel":
        if not payload.excel_bytes:
            raise ValueError("请选择 Excel 文件")
        return parse_excel(payload.excel_bytes)
    if payload.input_mode == "paste":
        return parse_paste(payload.pasted_text)
    raise ValueError(f"不支持的输入方式：{payload.input_mode}")
