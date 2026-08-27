from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class RosterWorkbook:
    def __init__(self, target: Path) -> None:
        self.target = target
        self._person_columns: list[str] = []
        self._relations: list[dict[str, str]] = []
        self.target.parent.mkdir(parents=True, exist_ok=True)

    @property
    def relations(self) -> list[dict[str, str]]:
        return [dict(row) for row in self._relations]

    def append_qualification(
        self,
        code: str,
        columns: list[str],
        people: list[dict[str, str]],
    ) -> list[dict[str, str]]:
        normalized_columns = [str(column or "").strip() for column in columns]
        if not normalized_columns or any(not column for column in normalized_columns):
            raise ValueError("门户人员字段不能为空")
        if self._person_columns and normalized_columns != self._person_columns:
            raise ValueError("不同运行资格返回的门户人员字段不一致")
        added = [
            {column: str(person.get(column, "") or "").strip() for column in normalized_columns}
            for person in people
        ]
        next_columns = self._person_columns or normalized_columns
        next_relations = [*self._relations, *({**person, "资质类别": code} for person in added)]
        self._save(next_columns, next_relations)
        self._person_columns = list(next_columns)
        self._relations = next_relations
        return added

    def _save(self, person_columns: list[str], relations: list[dict[str, str]]) -> None:
        headers = ["序号", *person_columns, "资质类别"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "运行资质人员名册"
        sheet.append(headers)
        for index, relation in enumerate(relations, start=1):
            sheet.append([index, *(relation.get(column, "") for column in person_columns), relation["资质类别"]])
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"
        width_by_header = {
            "序号": 10,
            "员工号": 16,
            "姓名": 16,
            "技术信息": 30,
            "出生日期": 14,
            "所属部门": 28,
            "资质类别": 16,
        }
        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width_by_header.get(header, 18)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
            cell.alignment = Alignment(horizontal="center")
        staff_column = headers.index("员工号") + 1
        for row in sheet.iter_rows(min_row=2, min_col=staff_column, max_col=staff_column):
            row[0].number_format = "@"
        temporary = self.target.with_name(f".{self.target.name}.tmp.xlsx")
        try:
            workbook.save(temporary)
            temporary.replace(self.target)
        finally:
            workbook.close()
            if temporary.exists():
                temporary.unlink()

    def read_headers(self) -> list[str]:
        workbook = load_workbook(self.target, read_only=True, data_only=True)
        try:
            return [str(value or "") for value in next(workbook.active.iter_rows(values_only=True), ())]
        finally:
            workbook.close()

    def read_rows(self) -> list[list[object]]:
        workbook = load_workbook(self.target, read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(min_row=2, values_only=True)]
        finally:
            workbook.close()
