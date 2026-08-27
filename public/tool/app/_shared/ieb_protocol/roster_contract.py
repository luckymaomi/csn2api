from __future__ import annotations

import io
import re
import time
import zipfile

from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook

from .base import IebContractError
from .roster_models import (
    FilterField,
    OperatingQualification,
    QueryFilters,
    SelectOption,
)


FORM_SELECTOR = "#showEmpProfileCompositeListPageForm"
FIELD_LABELS = {
    "primaryBaseArray": "注册基地",
    "baseArray": "运行基地",
    "techBase": "技术基地",
    "fleetCdToArray": "执飞机型",
    "bolMultiQualCd": "是否混飞",
    "qualTypeArray": "技术等级",
    "bolPriBase": "是否帮飞",
    "bolJCY": "检查员/教员",
    "activeStatusArray": "在职状态",
    "fleetCdbranch": "分部",
}
OOXML_REQUIRED_PARTS = {"[Content_Types].xml", "xl/workbook.xml"}
SELECTION_HEADERS = {"", "#"}
RESULT_RESERVED_HEADERS = {"序号", "资质类别"}


def _text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_filter_fields(html: str) -> list[FilterField]:
    soup = BeautifulSoup(str(html or ""), "html.parser")
    form = soup.select_one(FORM_SELECTOR)
    if form is None:
        raise IebContractError("技术资料页面没有返回有效筛选表单")
    fields: list[FilterField] = []
    for name, label in FIELD_LABELS.items():
        node = form.find("select", attrs={"name": name})
        if not isinstance(node, Tag):
            raise IebContractError(f"技术资料筛选表单缺少字段：{label}")
        options = [
            SelectOption(_text(option.get("value")), _text(option.get_text(" ", strip=True)))
            for option in node.select("option")
        ]
        if not options or any(not option.label for option in options):
            raise IebContractError(f"技术资料筛选字段没有有效选项：{label}")
        fields.append(FilterField(name, label, node.has_attr("multiple"), options))
    return fields


def parse_operating_qualifications(html: str) -> list[OperatingQualification]:
    soup = BeautifulSoup(str(html or ""), "html.parser")
    results: list[OperatingQualification] = []
    seen: set[str] = set()
    for code_cell in soup.select(".qualCdTd"):
        code = _text(code_cell.get("data-qualcd") or code_cell.get("data-qualCd"))
        row = code_cell.find_parent("tr")
        cells = row.find_all("td", recursive=False) if isinstance(row, Tag) else []
        values = [_text(cell.get_text(" ", strip=True)) for cell in cells]
        if not code or code in seen or len(values) < 6:
            raise IebContractError("运行资格动态列表存在缺失或重复代码")
        seen.add(code)
        results.append(OperatingQualification(code, values[2], values[3], values[4], values[5]))
    if not results:
        raise IebContractError("运行资格动态列表没有返回有效选项")
    return results


def build_filter_params(
    filters: QueryFilters,
    qualification_code: str,
    page: int | None = None,
    timestamp: int | None = None,
) -> list[tuple[str, str]]:
    code = _text(qualification_code)
    if not code:
        raise ValueError("运行资格代码不能为空")
    params: list[tuple[str, str]] = [
        ("personName", filters.person_name),
        ("staffNumAllDesc", ""),
        ("primaryBaseArray", filters.primary_base),
        ("baseArray", filters.operating_base),
        ("techBase", filters.technical_base),
    ]
    params.extend(("fleetCdToArray", value) for value in filters.fleet_codes)
    if not filters.fleet_codes:
        params.append(("fleetCdToArray", ""))
    params.append(("bolMultiQualCd", filters.multi_qualification))
    params.extend(("qualTypeArray", value) for value in filters.technical_levels)
    params.extend([
        ("bolPriBase", filters.help_flying),
        ("bolJCY", filters.instructor_checker),
    ])
    params.extend(("activeStatusArray", value) for value in filters.active_statuses)
    params.extend([
        ("fleetCdbranch", filters.branch),
        ("isOperQual", "Y"),
        ("operQualArray", code),
    ])
    if filters.deviation:
        params.append(("isPl", "Y"))
    if page is not None:
        params.append(("page", str(page)))
    params.append(("currentStr", str(timestamp or int(time.time() * 1000))))
    return params


def _table_headers(table: Tag) -> list[str]:
    headers = [_text(cell.get_text(" ", strip=True)) for cell in table.select("thead th")]
    if not headers:
        headers = [_text(cell.get_text(" ", strip=True)) for cell in table.select("tr th")]
    return headers


def parse_query_result(html: str) -> tuple[int, list[str], list[dict[str, str]]]:
    soup = BeautifulSoup(str(html or ""), "html.parser")
    text = _text(soup.get_text(" ", strip=True))
    match = re.search(r"共\s*(\d+)\s*条记录", text)
    if match:
        total = int(match.group(1))
    elif re.search(r"(?:没有(?:找到)?相关信息|暂无相关信息|无相关信息)", text):
        total = 0
    else:
        raise IebContractError("组合筛选响应没有返回可识别的记录总数")

    header_candidates = [_table_headers(table) for table in soup.select("table")]
    headers = next((item for item in header_candidates if "员工号" in item and "姓名" in item), [])
    rows: list[dict[str, str]] = []
    if headers:
        for table in soup.select("table"):
            for row in table.select("tbody tr"):
                cells = row.find_all("td", recursive=False)
                values = [_text(cell.get("title") or cell.get_text(" ", strip=True)) for cell in cells]
                if len(values) == len(headers) and any(values):
                    rows.append(dict(zip(headers, values)))
    return total, headers, rows


def parse_query_count(html: str) -> int:
    return parse_query_result(html)[0]


def personnel_columns(query_headers: list[str]) -> list[str]:
    columns = [_text(header) for header in query_headers if _text(header) not in SELECTION_HEADERS]
    if "员工号" not in columns or "姓名" not in columns:
        raise IebContractError("门户人员查询表缺少员工号或姓名表头")
    if len(set(columns)) != len(columns):
        raise IebContractError("门户人员查询表存在重复字段")
    reserved = [column for column in columns if column in RESULT_RESERVED_HEADERS]
    if reserved:
        raise IebContractError(f"门户人员查询表占用结果字段：{', '.join(reserved)}")
    return columns


def validate_ooxml(content: bytes) -> None:
    if not isinstance(content, bytes) or not content.startswith(b"PK"):
        raise IebContractError("门户普通导出未返回 OOXML 文件")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            if not OOXML_REQUIRED_PARTS.issubset(names) or not any(
                name.startswith("xl/worksheets/") and name.endswith(".xml") for name in names
            ):
                raise IebContractError("门户普通导出缺少 OOXML 必要部件")
    except IebContractError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise IebContractError("门户普通导出不是有效 OOXML 文件") from error


def _cell_text(value: object, identifier: bool = False) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if identifier and isinstance(value, int):
        return str(value).zfill(6)
    if identifier and isinstance(value, float) and value.is_integer():
        return str(int(value)).zfill(6)
    text = _text(value)
    return text.zfill(6) if identifier and text.isdigit() and len(text) <= 6 else text


def parse_ooxml_rows(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    validate_ooxml(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        source_rows = sheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(source_rows, ())]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or "员工号" not in headers or "姓名" not in headers:
            raise IebContractError("门户普通导出缺少员工号或姓名表头")
        if any(not value for value in headers) or len(set(headers)) != len(headers):
            raise IebContractError("门户普通导出表头存在空值或重复字段")
        staff_index = headers.index("员工号")
        rows: list[dict[str, str]] = []
        for source in source_rows:
            values = [
                _cell_text(source[index] if index < len(source) else None, index == staff_index)
                for index in range(len(headers))
            ]
            if any(values):
                rows.append(dict(zip(headers, values)))
        return headers, rows
    finally:
        workbook.close()


def project_personnel_rows(
    query_headers: list[str],
    export_headers: list[str],
    export_rows: list[dict[str, str]],
) -> tuple[list[str], list[dict[str, str]]]:
    columns = personnel_columns(query_headers)
    missing = [column for column in columns if column not in export_headers]
    if missing:
        raise IebContractError(f"门户普通导出缺少页面字段：{', '.join(missing)}")
    people = [{column: _text(row.get(column)) for column in columns} for row in export_rows]
    if any(not row["员工号"] or not row["姓名"] for row in people):
        raise IebContractError("门户普通导出存在员工号或姓名缺失的人员行")
    return columns, people
