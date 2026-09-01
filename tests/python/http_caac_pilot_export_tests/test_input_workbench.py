from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[3] / "public" / "tool" / "app" / "http-caac-pilot-export-helper"
sys.path.insert(0, str(ROOT))

from caac_pilot_export.input_data import read_excel_records  # noqa: E402
from caac_pilot_export.models import PilotOutcome, PilotRecord  # noqa: E402
from caac_pilot_export.portal_client import PortalClient, PortalContractError  # noqa: E402
from caac_pilot_export.runner import BatchRunner  # noqa: E402


def input_workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class InputWorkbenchTests(unittest.TestCase):
    def test_template_and_excel_input_use_name_with_optional_archive_number(self) -> None:
        template = load_workbook(ROOT / "template.xlsx", data_only=True)
        self.assertEqual([cell.value for cell in template.active[1]], ["姓名", "档案号"])
        self.assertEqual(template.active.max_row, 1)
        template.close()

        content = input_workbook([
            ["姓名", "档案号"],
            [" 张三 ", "00008277"],
            ["李四", ""],
            ["张三", "00008277"],
            ["", "00000001"],
        ])
        records, issues = read_excel_records(BytesIO(content))
        self.assertEqual([(item.name, item.archive_no) for item in records], [("张三", "00008277"), ("李四", "")])
        self.assertEqual(len(issues), 2)
        self.assertIn("重复", str(issues[0]))
        self.assertIn("姓名不能为空", str(issues[1]))

    def test_name_lookup_paginates_and_requires_unique_exact_ownership(self) -> None:
        requested = PilotRecord(0, {}, input_name="张三", input_archive_no="00000002", source="Excel 第2行")
        client = PortalClient()
        client.fetch_page = Mock(side_effect=[
            ([{"pilotId": "1", "pilotname": "张三", "archiveNo": "00000001"}], 2),
            ([{"pilotId": "2", "pilotname": "张三", "archiveNo": "00000002"}], 2),
        ])
        resolved = client.resolve_pilot(requested, page_size=1)
        self.assertEqual(resolved.pilot_id, "2")
        self.assertEqual(resolved.source, "Excel 第2行")
        self.assertEqual(client.fetch_page.call_count, 2)
        self.assertEqual(client.fetch_page.call_args_list[0].kwargs["pilotname"], "张三")

        ambiguous = PortalClient()
        ambiguous.fetch_page = Mock(return_value=([
            {"pilotId": "1", "pilotname": "张三", "archiveNo": "00000001"},
            {"pilotId": "2", "pilotname": "张三", "archiveNo": "00000002"},
        ], 2))
        with self.assertRaisesRegex(PortalContractError, "匹配到 2 人"):
            ambiguous.resolve_pilot(PilotRecord(0, {}, input_name="张三"))

    def test_runner_records_invalid_input_and_continues_without_request(self) -> None:
        invalid = PilotRecord(0, {}, source="Excel 第2行", input_error="姓名不能为空")
        valid = PilotRecord(1, {"pilotId": "2", "pilotname": "李四"}, source="Excel 第3行")
        portal = Mock()
        portal.query_one.return_value = PilotOutcome(valid, "成功", {"pilotId": "2"})
        result = BatchRunner(portal).run([invalid, valid], 2)
        self.assertEqual([item.status for item in result.outcomes], ["失败", "成功"])
        self.assertEqual(portal.query_one.call_count, 1)


if __name__ == "__main__":
    unittest.main()
