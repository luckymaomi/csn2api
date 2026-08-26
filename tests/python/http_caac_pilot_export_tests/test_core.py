from __future__ import annotations

import sys
import tempfile
import threading
import unittest
import urllib.request
from html.parser import HTMLParser
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3] / "public" / "tool" / "app" / "http-caac-pilot-export-helper"
sys.path.insert(0, str(ROOT))

from caac_pilot_export.credentials import parse_credentials  # noqa: E402
from caac_pilot_export.exporter import ResultExporter  # noqa: E402
from caac_pilot_export.models import PilotOutcome, PilotRecord  # noqa: E402
from caac_pilot_export.portal_client import PortalClient, PortalContractError, PortalSessionExpired  # noqa: E402
from caac_pilot_export.runner import BatchRunner  # noqa: E402
from caac_pilot_export.server import create_server  # noqa: E402


class _AssetParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.paths: list[str] = []

    def handle_starttag(self, _tag: str, attrs: list[tuple[str, str | None]]) -> None:
        for name, value in attrs:
            if name in {"href", "src"} and value and value.startswith("./"):
                self.paths.append("/" + value[2:])


class CoreTests(unittest.TestCase):
    def test_shell_cookie_escape_is_decoded(self) -> None:
        cookies = parse_credentials(r"-b $'yzzzm=a; JSESSIONID=abc\u0021xyz; home=h\_1'")
        self.assertEqual(cookies["JSESSIONID"], "abc!xyz")
        self.assertEqual(cookies["home"], "h_1")

    def test_page_parser_requires_total_and_rows(self) -> None:
        rows, total = PortalClient._parse_page({"total": 2, "rows": [{"pilotId": "1"}]})
        self.assertEqual((len(rows), total), (1, 2))
        with self.assertRaises(PortalContractError):
            PortalClient._parse_page({"rows": []})

    def test_runner_continues_after_one_detail_failure(self) -> None:
        records = [PilotRecord(i, {"pilotId": str(i + 1), "pilotname": f"P{i + 1}"}) for i in range(3)]
        portal = Mock()
        portal.query_one.side_effect = [RuntimeError("普通错误"), PilotOutcome(records[1], "成功", {"pilotId": "2"}), PilotOutcome(records[2], "成功", {"pilotId": "3"})]
        persisted = []
        result = BatchRunner(portal, persist=persisted.append).run(records, 3)
        self.assertEqual([item.status for item in result.outcomes], ["失败", "成功", "成功"])
        self.assertEqual(portal.query_one.call_count, 3)
        self.assertEqual([item.status for item in persisted], ["失败", "成功", "成功"])

    def test_runner_stops_after_session_expired(self) -> None:
        records = [PilotRecord(i, {"pilotId": str(i + 1)}) for i in range(2)]
        portal = Mock()
        portal.query_one.side_effect = PortalSessionExpired("expired")
        result = BatchRunner(portal).run(records, 2)
        self.assertEqual(result.outcomes[0].status, "失败")
        self.assertEqual(result.outcomes[1].status, "未查询")
        self.assertTrue(result.session_expired)
        self.assertEqual(portal.query_one.call_count, 1)

    def test_exporter_writes_one_sheet_and_dynamic_columns(self) -> None:
        outcome = PilotOutcome(PilotRecord(0, {"pilotId": "1", "pilotname": "示例"}), "成功", {"pilotId": "1", "pilotPhotoFileUrl": None, "layerInfosList": [{"layerinfo": "根", "sublayers": [{"layerinfo": "子", "sublayers": []}]}], "educationList": [{"school": "学校"}]})
        with tempfile.TemporaryDirectory() as directory:
            path = ResultExporter(Path(directory), "run").write([outcome])
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["飞行员信息"])
            headers = [cell.value for cell in workbook.active[1]]
            self.assertIn("执照等级", headers)
            self.assertIn("学习简历", headers)
            self.assertIn("照片文件URL", headers)
            self.assertIn("根\n  子", [cell.value for cell in workbook.active[2]])

    def test_exporter_saves_each_completed_pilot_as_checkpoint(self) -> None:
        first = PilotOutcome(PilotRecord(0, {"pilotId": "1"}), "成功", {"pilotId": "1", "firstField": "A"})
        second = PilotOutcome(PilotRecord(1, {"pilotId": "2"}), "成功", {"pilotId": "2", "laterField": "B"})
        with tempfile.TemporaryDirectory() as directory:
            exporter = ResultExporter(Path(directory), "checkpoint")
            path = exporter.append(first)
            first_book = load_workbook(path)
            self.assertEqual(first_book.active.max_row, 2)
            exporter.append(second)
            second_book = load_workbook(path)
            self.assertEqual(second_book.active.max_row, 3)
            self.assertIn("未映射字段（laterField）", [cell.value for cell in second_book.active[1]])

    def test_epoch_milliseconds_and_deep_qualification_tree_are_readable(self) -> None:
        from caac_pilot_export.flatten import flatten_detail
        flattened = flatten_detail({"birthday": 297014400000, "medicalBirthday": "1979-06-01", "licenseLimit": "第一条</br>第二条", "layerInfosList": [{"layerId": "1", "layerinfo": "根", "limitEnglish": "限制", "sublayers": [{"layerId": "2", "layerinfo": "子", "sublayers": [{"layerId": "3", "layerinfo": "孙", "sublayers": []}]}]}]})
        self.assertEqual(flattened["出生日期"], "1979-06-01")
        self.assertEqual(flattened["执照备注"], "第一条\n第二条")
        self.assertEqual(flattened["执照备注原文"], "第一条\n第二条")
        self.assertEqual(flattened["执照等级"], "根；等级ID：1；英语限制：限制\n  子；等级ID：2\n    孙；等级ID：3")

    def test_license_remarks_combine_every_page_source_field(self) -> None:
        from caac_pilot_export.flatten import flatten_detail

        flattened = flatten_detail({
            "englishLevel": "4",
            "validDate": 1831564800000,
            "chineseLevel": "6级",
            "chineseLevelValidDate": None,
            "licenseLimit": "第一条</br>第二条",
        })

        self.assertEqual(
            flattened["执照备注"],
            "英语语言能力 4级，无线电通信资格，有效期至 2028-01-16\n"
            "汉语语言能力 6级，无线电通信资格\n"
            "第一条\n第二条",
        )
        self.assertEqual(flattened["英语语言等级"], "4")
        self.assertEqual(flattened["英语语言等级有效期"], "2028-01-16")
        self.assertEqual(flattened["汉语语言等级"], "6级")
        self.assertEqual(flattened["执照备注原文"], "第一条\n第二条")

    def test_detail_cleaner_preserves_every_dynamic_record(self) -> None:
        from caac_pilot_export.flatten import flatten_detail

        records = [{"PCID": str(index), "CHECKDATE": 1778947200000} for index in range(1, 51)]
        text = flatten_detail({"checkpcinfos": records})["熟练检查"]

        self.assertEqual(len(text.splitlines()), 50)
        self.assertIn("50. 检查日期：2026-05-17；熟练检查ID：50", text)

    def test_detail_cleaner_translates_records_and_preserves_backend_fields(self) -> None:
        from caac_pilot_export.flatten import flatten_detail

        flattened = flatten_detail({
            "printList": [
                {
                    "printId": "internal-print-id",
                    "licenseId": "internal-license-id",
                    "printReason": "英语语言等级更新",
                    "agreeUser": "SYS",
                    "agreeDate": 1778947200000,
                    "appId": "internal-app-id",
                    "cardNo": "internal-card-no",
                    "notes": "internal-note",
                    "creationBy": "SYS",
                    "creationDate": 1778947200000,
                    "extraFlag": "保留",
                },
                {"printReason": "补发", "agreeDate": 1779033600000},
            ]
        })

        text = flattened["正式执照打印情况"]
        self.assertIn("打印原因：英语语言等级更新", text)
        self.assertIn("同意日期：2026-05-17", text)
        self.assertIn("同意人：SYS", text)
        self.assertIn("申请ID：internal-app-id", text)
        self.assertIn("卡号：internal-card-no", text)
        self.assertIn("备注：internal-note", text)
        self.assertIn("创建人：SYS", text)
        self.assertIn("打印记录ID：internal-print-id", text)
        self.assertIn("执照ID：internal-license-id", text)
        self.assertIn("未映射字段（extraFlag）：保留", text)
        self.assertIn("\n2. 同意日期：2026-05-18；打印原因：补发", text)
        self.assertNotIn("appId=", text)
        self.assertNotIn("printId=", text)

    def test_detail_cleaner_uses_confirmed_page_section_labels(self) -> None:
        from caac_pilot_export.flatten import flatten_detail

        flattened = flatten_detail({
            "address": "地址",
            "idType": "证件",
            "idCode": "编号",
            "licenseTypeName": "执照",
            "licenseLimit": "备注",
            "englishLevel": "4",
            "pilotPhotoFileUrl": "https://example.test/photo.jpg",
            "jzhxinfos": [{"CHECKLINE": "航段", "EXAMINER": "检查员", "CHECKDATE": 1778947200000, "LASTEXPDATE": 1810483200000, "RESULT": "1", "CHECKPLACE": "内部地点"}],
            "yjscinfos": [{"CHECKPLACE": "训练地点", "CHECKDATE": 1778947200000, "LASTEXPDATE": 1810483200000, "RESULT": "1", "CHECKLINE": "内部航段", "EXAMINER": "内部人员"}],
            "checkpcinfos": [{"CHECKDATE": 1778947200000, "CHECKTYPE": "PC-PIC", "ACMODEL": "机型", "EXAMINER": "考试员", "NOTES": "备注", "NEXTDATE": 1810483200000, "RESULT": "1"}],
            "printhis": [{"PRINTDATE": 1778947200000, "USERNAME": "打印人", "PRINTREASON": "打印原因"}],
        })

        self.assertEqual(flattened["永久住址"], "地址")
        self.assertEqual(flattened["身份证明名称"], "证件")
        self.assertEqual(flattened["身份证明编号"], "编号")
        self.assertEqual(flattened["所持执照"], "执照")
        self.assertEqual(flattened["执照备注"], "英语语言能力 4级，无线电通信资格\n备注")
        self.assertEqual(flattened["执照备注原文"], "备注")
        self.assertEqual(flattened["英语语言等级"], "4")
        self.assertEqual(flattened["照片文件URL"], "https://example.test/photo.jpg")
        self.assertIn("检查航段：航段", flattened["机长航线"])
        self.assertIn("检查状态：1", flattened["机长航线"])
        self.assertIn("检查地点：内部地点", flattened["机长航线"])
        self.assertIn("训练地点：训练地点", flattened["应急生存"])
        self.assertIn("训练日期：2026-05-17", flattened["应急生存"])
        self.assertIn("检查航段：内部航段", flattened["应急生存"])
        self.assertIn("检查员：内部人员", flattened["应急生存"])
        self.assertIn("考试员：考试员", flattened["熟练检查"])
        self.assertIn("下次检查日期：2027-05-17", flattened["熟练检查"])
        self.assertIn("检查状态：1", flattened["熟练检查"])
        self.assertIn("打印人：打印人", flattened["正式执照打印历史"])

    def test_confirmed_response_schema_has_a_chinese_mapping_for_every_field(self) -> None:
        from caac_pilot_export.field_mappings import RECORD_SPECS, TOP_LEVEL_FIELD_LABELS, TREE_FIELD_LABELS

        top_fields = {
            "address", "archiveNo", "birthday", "birthplace", "chineseLevel", "chineseLevelValidDate",
            "chineseOrg", "country", "degree", "englishLevel", "idCode", "idType", "licenseLimit",
            "licenseNo", "licenseStatus", "licenseTypeName", "mailaddr", "medicalBirthday", "medicalno",
            "mobile", "nation", "pilotDegreeFileUrl", "pilotId", "pilotIdFileUrl", "pilotPhotoFileUrl",
            "pilotgrade", "pilotname", "pilotnameeng", "regionName", "sex", "validDate", "zipcode",
            "nameCn", "expDate", "enddate", "copType",
        }
        record_fields = {
            "educationList": {"college", "creationBy", "creationDate", "degree", "eduId", "eindex", "endDate", "major", "note", "pilotId", "school", "startDate", "updateBy", "updateDate", "validFlag"},
            "workexps": {"creationBy", "creationDate", "eindex", "endDate", "location", "note", "organization", "pilotId", "startDate", "title", "updateBy", "updateDate", "validFlag", "workExpId"},
            "jzhxinfos": {"CHECKDATE", "CHECKLINE", "CHECKPLACE", "CHECK_ID", "EXAMINER", "LASTEXPDATE", "RESULT"},
            "yjscinfos": {"CHECKDATE", "CHECKLINE", "CHECKPLACE", "CHECK_ID", "EXAMINER", "LASTEXPDATE", "RESULT"},
            "printhis": {"PRINTDATE", "PRINTREASON", "USERNAME"},
            "printList": {"agreeDate", "agreeUser", "appId", "cardId", "cardNo", "creationBy", "creationDate", "licenseId", "notes", "printDate", "printId", "printReason", "printStatus", "printUser", "receiveFlag", "updateBy", "updateDate", "validDate", "validUser"},
            "checkpcinfos": {"ACMODEL", "ACMODELID", "BASEMONTH", "CHECKDATE", "CHECKTYPE", "EXAMINER", "LASTEXPDATE", "NEXTDATE", "NEXTEXPDATE", "NEXT_MONTH", "NOTES", "PCID", "RESULT"},
        }

        self.assertFalse(top_fields - set(TOP_LEVEL_FIELD_LABELS))
        self.assertEqual(TOP_LEVEL_FIELD_LABELS["licenseTypeName"], "所持执照")
        self.assertEqual(TOP_LEVEL_FIELD_LABELS["nameCn"], "所持执照")
        self.assertEqual(TOP_LEVEL_FIELD_LABELS["licenseLimit"], "执照备注原文")
        self.assertFalse({"layerId", "layerinfo", "limitEN", "limitEnglish", "limitZN"} - set(TREE_FIELD_LABELS))
        for section, fields in record_fields.items():
            with self.subTest(section=section):
                self.assertFalse(fields - set(RECORD_SPECS[section].fields))

    def test_detail_values_override_listing_values_without_duplicate_raw_headers(self) -> None:
        outcome = PilotOutcome(
            PilotRecord(0, {"pilotId": "1", "pilotname": "列表姓名", "listOnly": "列表值"}),
            "成功",
            {"pilotId": "1", "pilotname": "详情姓名", "detailOnly": "详情值"},
        )
        with tempfile.TemporaryDirectory() as directory:
            path = ResultExporter(Path(directory), "merge").write([outcome])
            workbook = load_workbook(path)
            headers = [cell.value for cell in workbook.active[1]]
            values = [cell.value for cell in workbook.active[2]]
            row = dict(zip(headers, values))
            self.assertEqual(row["姓名"], "详情姓名")
            self.assertEqual(row["未映射字段（listOnly）"], "列表值")
            self.assertEqual(row["未映射字段（detailOnly）"], "详情值")
            self.assertFalse(any(str(header).startswith(("列表.", "详情.")) for header in headers))

    def test_workbench_serves_every_local_ui_asset(self) -> None:
        manager = Mock()
        manager.snapshot.return_value = {"phase": "idle"}
        manager.download_path.return_value = None
        server = create_server(ROOT, 0, manager)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        host, port = server.server_address
        try:
            with urllib.request.urlopen(f"http://{host}:{port}/", timeout=5) as response:
                html = response.read().decode("utf-8")
            parser = _AssetParser()
            parser.feed(html)
            self.assertIn("/libs/bootstrap.min.css", parser.paths)
            self.assertIn('id="progressText"', html)
            self.assertIn('id="progressBar"', html)
            self.assertIn('role="progressbar"', html)
            self.assertIn('aria-valuenow="0"', html)
            self.assertIn('data-keepalive-widget', html)
            self.assertIn('id="keepAliveCountdown"', html)
            for path in parser.paths:
                with self.subTest(path=path):
                    with urllib.request.urlopen(f"http://{host}:{port}{path}", timeout=5) as response:
                        content_type = response.headers.get_content_type()
                        body = response.read()
                    self.assertGreater(len(body), 20)
                    if path.endswith(".css"):
                        self.assertEqual(content_type, "text/css")
                    elif path.endswith((".js", ".mjs")):
                        self.assertIn(content_type, {"text/javascript", "application/javascript"})
        finally:
            server.shutdown()
            thread.join(timeout=5)
            server.server_close()


if __name__ == "__main__":
    unittest.main()
