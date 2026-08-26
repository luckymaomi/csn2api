from __future__ import annotations

import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[3] / "public" / "tool" / "app" / "http-caac-pilot-export-helper"
sys.path.insert(0, str(ROOT))

from caac_pilot_export.manager import RunManager  # noqa: E402
from caac_pilot_export.models import InputPayload, PilotOutcome, PilotRecord  # noqa: E402


class FakeKeepAlive:
    def status(self, ready: bool) -> dict[str, object]:
        return {"state": "healthy" if ready else "stopped", "running": ready, "nextIntervalSeconds": 45}

    def stop(self) -> None:
        return

    def record_success(self) -> None:
        return

    def start(self, *_args: object) -> None:
        return


class FakeClient:
    def __init__(self) -> None:
        self.keepalive = FakeKeepAlive()
        self.copyable_credentials = "session=test"
        self.queries: list[PilotRecord] = []
        self.full_list_calls = 0

    def load_credentials(self, _text: str) -> dict[str, object]:
        return {"verifiedAt": "2026-08-26 09:00:00", "cookieCount": 1, "total": 1}

    def check_session(self) -> None:
        return

    def require_credentials(self) -> None:
        return

    def fetch_all_pilots(self) -> tuple[list[PilotRecord], int]:
        self.full_list_calls += 1
        return [PilotRecord(0, {"pilotId": "1", "pilotname": "全量示例", "archiveNo": "00000001"})], 1

    def query_one(self, record: PilotRecord) -> PilotOutcome:
        self.queries.append(record)
        resolved = record if record.pilot_id else PilotRecord(
            record.index,
            {"pilotId": "2", "pilotname": record.input_name, "archiveNo": record.input_archive_no},
            record.input_name,
            record.input_archive_no,
            record.source,
            "",
            record.input_row,
        )
        return PilotOutcome(resolved, "成功", {"pilotId": resolved.pilot_id, "pilotname": resolved.name, "dynamicField": "动态值"})

    def invalidate_credentials(self) -> None:
        return

    def close(self) -> None:
        return


def input_workbook(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class ManagerTests(unittest.TestCase):
    def test_selected_mode_persists_each_row_and_exposes_the_same_excel_data(self) -> None:
        client = FakeClient()
        content = input_workbook([
            ["姓名", "档案号"],
            ["名单示例", "00000002"],
            ["名单示例", "00000002"],
        ])
        payload = InputPayload("名单.xlsx", content)
        with tempfile.TemporaryDirectory() as directory:
            manager = RunManager(Path(directory), client)  # type: ignore[arg-type]
            manager.verify_credentials("session=test")
            checked = manager.check_data("selected", payload)
            self.assertEqual(checked["checks"]["data"]["validCount"], 1)  # type: ignore[index]
            self.assertEqual(checked["checks"]["data"]["invalidCount"], 1)  # type: ignore[index]
            manager.run("selected", payload)
            manager._thread.join(timeout=5)  # type: ignore[union-attr]

            state = manager.snapshot()
            self.assertEqual(state["phase"], "completed")
            self.assertEqual(state["progress"]["success"], 1)  # type: ignore[index]
            self.assertEqual(state["progress"]["failed"], 1)  # type: ignore[index]
            self.assertEqual(len(client.queries), 1)
            preview = manager.results_page(0, 50)
            self.assertEqual(preview["total"], 2)
            self.assertIn("未映射字段（dynamicField）", preview["columns"])

            result_path = manager.download_path()
            self.assertIsNotNone(result_path)
            workbook = load_workbook(result_path, data_only=True)
            headers = [cell.value for cell in workbook.active[1]]
            rows = [dict(zip(headers, [cell.value for cell in row])) for row in workbook.active.iter_rows(min_row=2)]
            self.assertEqual(rows[0]["未映射字段（dynamicField）"], preview["rows"][0]["未映射字段（dynamicField）"])
            self.assertEqual([row["抓取状态"] for row in rows], ["成功", "失败"])
            workbook.close()
            manager.shutdown()

    def test_all_mode_keeps_dynamic_list_discovery(self) -> None:
        client = FakeClient()
        with tempfile.TemporaryDirectory() as directory:
            manager = RunManager(Path(directory), client)  # type: ignore[arg-type]
            manager.verify_credentials("session=test")
            manager.run("all")
            manager._thread.join(timeout=5)  # type: ignore[union-attr]
            self.assertEqual(client.full_list_calls, 1)
            self.assertEqual(manager.snapshot()["progress"]["success"], 1)  # type: ignore[index]
            manager.shutdown()


if __name__ == "__main__":
    unittest.main()
