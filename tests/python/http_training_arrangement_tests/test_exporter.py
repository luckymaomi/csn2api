from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .common import SHARED_ROOT  # noqa: F401
from ieb_protocol.training_arrangement_models import TrainingWindow
from http_training_arrangement.exporter import ResultExporter
from http_training_arrangement.models import AreaWindow


class ExporterTests(unittest.TestCase):
    def test_writes_summary_and_one_sheet_per_period(self) -> None:
        window = TrainingWindow(
            "year", "年", date(2026, 1, 1), date(2026, 9, 10),
            ["训练大纲", "训练科目", "员工号/姓名"],
            [{"训练大纲": "777", "训练科目": "机长等级训练C转D", "员工号/姓名": "100001 (测试)"}],
            1, 1,
        )
        with tempfile.TemporaryDirectory() as directory:
            path = ResultExporter(Path(directory), "test-run").write([window]).excel
            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["汇总", "年"])
                self.assertEqual(workbook["汇总"].cell(2, 5).value, 1)
                self.assertEqual(workbook["年"].cell(2, 2).value, "机长等级训练C转D")
            finally:
                workbook.close()

    def test_writes_area_summary_and_details(self) -> None:
        window = AreaWindow(
            "区域航线", "week", "周", date(2026, 9, 3), date(2026, 9, 10),
            ["区域类型", "训练科目", "训练状态"],
            [{"区域类型": "区域航线", "训练科目": "特殊区域、航路和机场训练", "训练状态": "完成"}],
            1,
            1,
        )
        summary = [{
            "type": "区域航线",
            "week": {"count": 1, "passed": 1, "rate": 100.0},
            "month": {"count": 0, "passed": 0, "rate": None},
            "year": {"count": 1, "passed": 1, "rate": 100.0},
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = ResultExporter(Path(directory), "test-area").write_area([window], summary).excel
            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["分类汇总", "明细"])
                self.assertEqual([cell.value for cell in workbook["分类汇总"][1]], ["类型", "本周", None, "本月", None, "年度", None])
                self.assertEqual(workbook["分类汇总"].cell(3, 3).value, "100%")
                self.assertEqual(workbook["明细"].cell(2, 1).value, "区域航线")
                self.assertEqual(workbook["明细"].cell(2, 2).value, "周")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
