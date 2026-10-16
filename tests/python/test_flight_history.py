from __future__ import annotations

import sys
import unittest
import threading
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'public/tool/app/_shared'))
sys.path.insert(0, str(ROOT / 'public/tool/app/http-flight-history-helper'))

from ieb_protocol.base import IebContractError
from ieb_protocol.flight_history_contract import parse_history
from ieb_protocol.flight_detail import FlightDetailProtocol
from http_flight_history.input_data import parse_pasted_records, read_excel_records
from http_flight_history.models import QueryOutcome, QueryRecord, TableResult
from http_flight_history.runner import BatchRunner
from http_flight_history.exporter import ResultExporter
from http_flight_history.portal_client import PortalSessionExpired


HTML = '''<button onclick="showflyTimeDetail('1','900001','示例甲','2026','','777')">2026</button>
<table><tr><th>日期</th><th>执飞机型</th><th>飞行经历</th><th>飞行时间</th><th>新增字段</th></tr></table>
<table><tr><td>2026</td><td>777</td><td>10:12</td><td>12:34</td><td>年度</td></tr>
<tr><td>汇总</td><td></td><td>123:45</td><td title="456:07">456...</td><td>保留</td></tr></table>'''


class HistoryContractTests(unittest.TestCase):
    def test_keeps_all_rows_with_dynamic_headers_original_order_and_identity(self):
        headers, rows = parse_history(HTML, '900001', '示例甲')
        self.assertEqual(headers, ['日期', '执飞机型', '飞行经历', '飞行时间', '新增字段'])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['日期'], '2026')
        self.assertEqual(rows[1]['飞行时间'], '456:07')
        self.assertEqual(rows[1]['新增字段'], '保留')
        self.assertEqual(rows[1]['日期'], '汇总')

    def test_rejects_wrong_identity_missing_duplicate_and_width(self):
        for html, staff, name in [
            (HTML, '900002', '示例甲'), (HTML, '900001', '示例乙'),
            (HTML.replace('汇总', '合并'), '900001', '示例甲'),
            (HTML.replace('<td>保留</td>', ''), '900001', '示例甲'),
            (HTML.replace('<td>汇总</td>', '<td>汇总</td><td>汇总</td>'), '900001', '示例甲'),
            (HTML.replace('<th>新增字段</th>', '<th>飞行时间</th>'), '900001', '示例甲'),
        ]:
            with self.subTest(staff=staff, name=name, html=html), self.assertRaises(IebContractError):
                parse_history(html, staff, name)

    def test_rejects_invalid_year_row_duplicate_summary_and_mixed_identity(self):
        for html in [
            HTML.replace('<td>年度</td>', ''),
            HTML.replace('</table>', '<tr><td>汇总</td></tr></table>', 1),
            HTML + '''<button onclick="showflyTimeDetail('1','900002','示例乙','2026','','777')">2026</button>''',
        ]:
            with self.assertRaises(IebContractError):
                parse_history(html, '900001', '示例甲')

    def test_preserves_confirmed_external_airline_note_on_summary_label(self):
        html = HTML.replace('<td>汇总</td>', '<td><div>汇总\n( <a>外航经历</a> )</div></td>')
        headers, rows = parse_history(html, '900001', '示例甲')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]['日期'], '汇总 ( 外航经历 )')
        self.assertEqual(rows[1]['飞行时间'], '456:07')

    def test_request_and_catalog_share_confirmed_form_contract(self):
        import json
        session = Mock(base_url='https://ieb.csair.com')
        session.request.return_value.text = HTML
        protocol = FlightDetailProtocol(session)
        _, headers, rows = protocol.request_history('900001', '示例甲')
        catalog = json.loads((ROOT / 'public/tool/app/api-docs/catalog/flight-detail.json').read_text(encoding='utf-8'))
        endpoint = next(item for item in catalog['endpoints'] if item['id'] == 'query-history')
        args, kwargs = session.request.call_args
        self.assertEqual(args, (endpoint['method'], endpoint['path']))
        self.assertEqual(set(kwargs['data']), {p['name'] for p in endpoint['parameters'] if p['in'] == 'form'})
        self.assertEqual(kwargs['data']['staffNum'], '900001')
        self.assertEqual(len(rows), 2)


class HistoryWorkbenchTests(unittest.TestCase):
    def test_two_column_inputs_aliases_duplicates_and_invalid_rows(self):
        records, errors = parse_pasted_records('员工号\t姓名\n900001\t示例甲\n900001 示例甲\n900002\n900003 示例丙')
        self.assertEqual([r.employee_id for r in records], ['900001', '900003'])
        self.assertEqual(len(errors), 2)
        book = Workbook()
        book.active.append(['员工姓名', '工号'])
        book.active.append(['示例甲', 12.0])
        book.active.append([None, 900002])
        buffer = BytesIO()
        book.save(buffer)
        book.close()
        records, errors = read_excel_records(buffer.getvalue())
        self.assertEqual((records[0].employee_id, records[0].name), ('000012', '示例甲'))
        self.assertEqual(len(errors), 1)

    def test_export_preserves_every_row_empty_values_and_duration_text(self):
        headers, rows = parse_history(HTML, '900001', '示例甲')
        record = QueryRecord('900001', '示例甲', '测试行')
        outcome = QueryOutcome(0, record, '成功', TableResult(headers, rows))
        with TemporaryDirectory() as directory:
            path = ResultExporter(Path(directory), 'test').write([record], [], [outcome]).original
            book = load_workbook(path, data_only=True)
            try:
                sheet = book.worksheets[0]
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual([c.value for c in sheet[1]], ['员工号', '姓名', *headers])
                self.assertEqual(sheet.cell(3, 6).value, '456:07')
                self.assertEqual(sheet.cell(3, 4).value, None)
                self.assertEqual(sheet.cell(3, 6).data_type, 's')
            finally:
                book.close()

    def test_serial_failures_continue_session_expiry_and_stop_keep_checkpoints(self):
        records = [QueryRecord(f'90000{i}', f'示例{i}', str(i)) for i in range(1, 4)]
        portal = Mock()
        portal.query.side_effect = [ValueError('行级异常'), TableResult(['日期'], [{'日期': '汇总'}]), PortalSessionExpired('会话失效')]
        checkpoints = []
        result = BatchRunner(portal).run(records, lambda outcomes: checkpoints.append(len(outcomes)))
        self.assertEqual((result.success, result.failed, result.session_expired), (1, 2, True))
        self.assertEqual(checkpoints, [1, 2, 3])
        self.assertEqual([call.args[0] for call in portal.query.call_args_list], records)
        stop = threading.Event()
        def first(_record):
            stop.set()
            return TableResult(['日期'], [{'日期': '汇总'}])
        portal.query.side_effect = first
        portal.query.reset_mock()
        result = BatchRunner(portal, stop).run(records)
        self.assertTrue(result.stopped)
        self.assertEqual(portal.query.call_count, 1)
        self.assertEqual(len(result.outcomes), 3)


if __name__ == '__main__':
    unittest.main()
