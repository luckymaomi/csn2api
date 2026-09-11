from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
APP_ROOT = ROOT / "public" / "tool" / "app"


class TemplateAssetTests(unittest.TestCase):
    def test_lock_template_matches_standard_workbook_and_keeps_type_dropdown(self) -> None:
        standard = ROOT / "public" / "template" / "通用批量锁班模板.xlsx"
        delivered = APP_ROOT / "http-lock-entry-helper" / "template.xlsx"
        self.assertEqual(delivered.read_bytes(), standard.read_bytes())

        workbook = load_workbook(delivered, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, ["锁班数据", "类型列表"])
            sheet = workbook.active
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["员工号", "姓名", "锁班类型", "开始日期", "结束日期"],
            )
            self.assertEqual(sheet.cell(2, 1).value, "000000")
            self.assertEqual(len(sheet.data_validations.dataValidation), 1)
            validation = sheet.data_validations.dataValidation[0]
            self.assertEqual(validation.type, "list")
            self.assertEqual(validation.formula1, "类型列表!$A$1:$A$51")
            self.assertEqual(str(validation.sqref), "C2:C1000")
            self.assertEqual(workbook["类型列表"].max_row, 51)
        finally:
            workbook.close()

    def test_non_dropdown_templates_have_one_clear_example_row(self) -> None:
        expected = {
            "http-caac-pilot-export-helper": (
                ["姓名", "档案号"],
                ["示例姓名", "00000000"],
            ),
            "http-flight-detail-helper": (
                ["员工号", "姓名", "开始年月", "结束年月"],
                ["000000", "示例姓名", "2026/01", "2026/12"],
            ),
            "http-flight-stats-helper": (
                ["员工号", "姓名", "开始日期", "结束日期"],
                ["000000", "示例姓名", "2026-01-01", "2026-01-31"],
            ),
            "http-qualification-query-helper": (
                ["员工号", "姓名"],
                ["000000", "示例姓名"],
            ),
        }
        for app_name, (headers, example) in expected.items():
            with self.subTest(app=app_name):
                workbook = load_workbook(APP_ROOT / app_name / "template.xlsx", data_only=True)
                try:
                    sheet = workbook.active
                    self.assertEqual([cell.value for cell in sheet[1]], headers)
                    self.assertEqual([cell.value for cell in sheet[2]], example)
                finally:
                    workbook.close()

    def test_all_project_templates_have_no_colored_fill_or_font(self) -> None:
        template_paths = list(APP_ROOT.glob("*/template.xlsx"))
        template_paths.extend((ROOT / "public" / "template").glob("*.xlsx"))
        for path in sorted(template_paths):
            with self.subTest(template=str(path.relative_to(ROOT))):
                workbook = load_workbook(path, data_only=False)
                try:
                    for sheet in workbook.worksheets:
                        for row in sheet.iter_rows():
                            for cell in row:
                                self.assertIsNone(cell.fill.fill_type, f"{path}:{sheet.title}!{cell.coordinate}")
                                color = cell.font.color
                                if color is None:
                                    continue
                                self.assertIn(color.type, {"theme", "indexed", "rgb"})
                                if color.type == "theme":
                                    self.assertEqual(color.theme, 1)
                                elif color.type == "rgb":
                                    self.assertIn(color.rgb, {"00000000", "FF000000"})
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
