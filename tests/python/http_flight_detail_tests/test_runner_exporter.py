from __future__ import annotations

import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import load_workbook

from .common import APP_ROOT  # noqa: F401
from http_flight_detail.exporter import DETAIL_OUTPUT_HEADERS, ResultExporter
from http_flight_detail.models import QueryRecord
from http_flight_detail.portal_client import PortalError, PortalSessionExpired
from http_flight_detail.runner import BatchRunner


def record(employee_id: str, start: str, end: str) -> QueryRecord:
    return QueryRecord(employee_id, f"示例{employee_id[-1]}", start, end, "测试")


class SerialPortal:
    def __init__(self) -> None:
        self.calls: list[tuple[str, str]] = []
        self.active = 0
        self.max_active = 0

    def load_timeline(self, item: QueryRecord):
        return [{"fleet": "32A"}]

    def query_month(self, item: QueryRecord, month: str, _timeline):
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        self.calls.append((item.employee_id, month))
        try:
            return DETAIL_OUTPUT_HEADERS[:1] + ["航班号", "机型", "日期", "左座", "右座", "操纵", "观察", "离场", "到达", "飞行时间", "夜航", "飞行阶段", "人工飞行时间", "人工飞行方式"], [{
                "序号": "1", "航班号": "CZ3001", "机型": "32A", "日期": f"{month}-01",
                "左座": "1", "右座": "", "操纵": "1", "观察": "", "离场": "CAN", "到达": "PEK",
                "飞行时间": "03:10", "夜航": "01:00", "飞行阶段": "巡航",
                "人工飞行时间": "00:12", "人工飞行方式": "自动驾驶",
            }]
        finally:
            self.active -= 1


class FailingPortal(SerialPortal):
    def query_month(self, item: QueryRecord, month: str, timeline):
        if month == "2026-02":
            self.calls.append((item.employee_id, month))
            raise PortalError("单月查询失败")
        return super().query_month(item, month, timeline)


class ExpiringPortal(SerialPortal):
    def query_month(self, item: QueryRecord, month: str, timeline):
        if month == "2026-02":
            self.calls.append((item.employee_id, month))
            raise PortalSessionExpired("登录凭据已失效")
        return super().query_month(item, month, timeline)


class RunnerExporterTests(unittest.TestCase):
    def test_queries_people_and_months_strictly_serial(self) -> None:
        items = [record("900001", "2026-01", "2026-02"), record("900002", "2026-03", "2026-03")]
        portal = SerialPortal()

        result = BatchRunner(portal).run(items)

        self.assertEqual(portal.max_active, 1)
        self.assertEqual(portal.calls, [
            ("900001", "2026-01"),
            ("900001", "2026-02"),
            ("900002", "2026-03"),
        ])
        self.assertEqual([item.status for item in result.outcomes], ["成功", "成功", "成功"])

    def test_single_month_failure_continues_but_expiry_stops_and_marks_remaining(self) -> None:
        items = [record("900001", "2026-01", "2026-03")]
        failed = BatchRunner(FailingPortal()).run(items)
        self.assertEqual([item.status for item in failed.outcomes], ["成功", "失败", "成功"])

        expired_portal = ExpiringPortal()
        expired = BatchRunner(expired_portal).run(items)
        self.assertEqual(expired_portal.calls, [("900001", "2026-01"), ("900001", "2026-02")])
        self.assertTrue(expired.session_expired)
        self.assertEqual([item.status for item in expired.outcomes], ["成功", "失败", "未查询"])

    def test_stop_marks_remaining_months_without_requesting_them(self) -> None:
        stop = threading.Event()
        portal = SerialPortal()

        def emit(event):
            if event.get("type") == "month_result":
                stop.set()

        result = BatchRunner(portal, stop_event=stop, emit=emit).run([
            record("900001", "2026-01", "2026-03")
        ])

        self.assertEqual(portal.calls, [("900001", "2026-01")])
        self.assertEqual([item.status for item in result.outcomes], ["成功", "未查询", "未查询"])

    def test_checkpoint_export_contains_detail_and_month_status(self) -> None:
        items = [record("900001", "2026-01", "2026-02")]
        result = BatchRunner(SerialPortal()).run(items)

        with tempfile.TemporaryDirectory() as directory:
            exporter = ResultExporter(Path(directory), "test-run")
            path = exporter.write(items, ["第3行：月份无效"], result.outcomes[:1])
            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["航班明细", "查询状态", "输入错误"])
                detail_headers = [cell.value for cell in workbook["航班明细"][1]]
                self.assertEqual(detail_headers, DETAIL_OUTPUT_HEADERS)
                self.assertEqual(workbook["航班明细"]["B2"].value, "示例1")
                self.assertEqual(workbook["航班明细"]["C2"].value, "900001")
                statuses = [workbook["查询状态"].cell(row, 5).value for row in (2, 3)]
                self.assertEqual(statuses, ["成功", "未查询"])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
