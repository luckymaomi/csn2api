import unittest
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

APP_ROOT = Path(__file__).resolve().parents[3] / "public" / "tool" / "app" / "http-soc-positioning-helper"
SHARED_ROOT = Path(__file__).resolve().parents[3] / "public" / "tool" / "app" / "_shared"
TEMPLATE_PATH = APP_ROOT / "template.xlsx"
sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(SHARED_ROOT))

from soc_protocol import APPLY_REASON_LABELS  # noqa: E402
from soc_positioning.input_data import normalize_date, parse_excel  # noqa: E402


class SocInputTests(unittest.TestCase):
    def test_delivery_template_leaves_optional_columns_empty(self):
        workbook = load_workbook(TEMPLATE_PATH, data_only=False, read_only=True)
        try:
            sheet = workbook["随机置位申请"]
            self.assertEqual(sheet.cell(1, 1).value, "序号")
            self.assertEqual(sheet.cell(1, 10).value, "是否需要录入")
            self.assertIsNone(sheet.cell(2, 1).value)
            self.assertIsNone(sheet.cell(2, 10).value)
            self.assertEqual(workbook.properties.title, "随机置位录入模板")
        finally:
            workbook.close()

    def test_date_normalization_handles_excel_serial_and_explicit_text(self):
        self.assertEqual(normalize_date(46037), "2026-01-15")
        self.assertEqual(normalize_date("2026/1/15"), "2026-01-15")
        self.assertEqual(normalize_date(date(2026, 1, 15)), "2026-01-15")

    def test_parse_excel_uses_reason_dictionary_and_skips_no(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "员工号", "姓名", "航班日期", "出发地", "目的地", "航班号", "申请原因", "备注", "是否需要录入"])
        sheet.append([1, "000001", "测试员工", 46037, "广州", "长沙", "CZ1234", "学习培训", "示例备注", "是"])
        sheet.append([2, "000002", "跳过员工", 46037, "广州", "郑州", "CZ5678", "不存在", "", "否"])
        stream = BytesIO(); workbook.save(stream); workbook.close()
        records, errors = parse_excel(stream.getvalue())
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["申请原因代码"], "6")
        self.assertEqual(records[0]["航班日期"], "2026-01-15")
        self.assertFalse(errors)
        self.assertIn("6", APPLY_REASON_LABELS)


if __name__ == "__main__":
    unittest.main()
