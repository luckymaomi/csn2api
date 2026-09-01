from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "public" / "tool" / "app" / "http-flight-detail-helper"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT.parent / "_shared"))

from http_flight_detail.exporter import DETAIL_OUTPUT_HEADERS, ResultExporter  # noqa: E402
from http_flight_detail.manager import RunManager  # noqa: E402
from http_flight_detail.models import MonthOutcome, QueryRecord  # noqa: E402


class PressureTests(unittest.TestCase):
    """手动压力门禁，不纳入常规 test:python。"""

    def test_9000_results_snapshot_page_and_append_export(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            class Client:
                def close(self):
                    return None

            manager = RunManager(Path(directory), client=Client())
            row = {header: "值" for header in DETAIL_OUTPUT_HEADERS[3:]}
            started = perf_counter()
            for index in range(9000):
                manager._handle_event({
                    "type": "month_result", "index": index, "recordIndex": index,
                    "employeeId": str(900000 + index), "name": f"示例{index}", "month": "2026-01",
                    "status": "成功", "headers": [], "rows": [row], "error": "",
                })
            snapshot = manager.snapshot()
            snapshot_json = json.dumps(snapshot, ensure_ascii=False)
            self.assertNotIn("results", snapshot)
            self.assertNotIn("航班号", snapshot_json)
            self.assertLess(len(snapshot_json), 200_000)
            status_page = manager.results_page("status", 8950, 100)
            detail_page = manager.results_page("detail", 8950, 100)
            self.assertEqual((status_page["total"], len(status_page["rows"])), (9000, 50))
            self.assertEqual((detail_page["total"], len(detail_page["rows"])), (9000, 50))
            self.assertLess(perf_counter() - started, 5.0)
            manager.shutdown()

    def test_9000_row_append_is_incremental_and_final_excel_is_streamed(self) -> None:
        records = [QueryRecord(str(900000 + index), f"示例{index}", "2026-01", "2026-01", "压力") for index in range(9000)]
        row = {header: "值" for header in DETAIL_OUTPUT_HEADERS[3:]}
        outcomes = [MonthOutcome(index, index, record, "2026-01", "成功", [row], DETAIL_OUTPUT_HEADERS[3:], "") for index, record in enumerate(records)]
        with tempfile.TemporaryDirectory() as directory:
            exporter = ResultExporter(Path(directory), "pressure")
            for end in range(1000, 9001, 1000):
                exporter.append_batch(outcomes[:end])
                exporter.flush()
            self.assertFalse(exporter.path.exists())
            self.assertTrue((Path(directory) / "detail-rows.jsonl").is_file())
            exporter.finalize(records, [])
            workbook = load_workbook(exporter.path, read_only=True)
            try:
                self.assertEqual(sum(1 for _ in workbook["航班明细"].iter_rows()), 9001)
            finally:
                workbook.close()
            exporter.close()
            self.assertFalse((Path(directory) / "detail-rows.jsonl").exists())


if __name__ == "__main__":
    unittest.main()
